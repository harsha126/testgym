import csv
import openpyxl
from datetime import datetime, date


def format_date(value):
    """Convert datetime/date/string to YYYY-MM-DD string, or return None if invalid."""
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        try:
            d = value.date() if isinstance(value, datetime) else value
            if date(2000, 1, 1) <= d <= date(2035, 12, 31):
                return d.strftime("%Y-%m-%d")
        except Exception:
            pass
    # Try parsing string
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
            try:
                d = datetime.strptime(value.strip(), fmt).date()
                if date(2000, 1, 1) <= d <= date(2035, 12, 31):
                    return d.strftime("%Y-%m-%d")
            except ValueError:
                pass
    return None


VALID_MODES = {
    "monthly": "Monthly",
    "quarterly": "Quarterly",
    "quaterly": "Quarterly",  # common misspelling
}


def normalize_mode(value):
    if value is None:
        return None
    return VALID_MODES.get(str(value).strip().lower())


def is_valid_amount(value):
    try:
        amount = float(value)
        return amount > 0
    except (ValueError, TypeError):
        return False


wb = openpyxl.load_workbook("BOLLARAM FEES LIST (1).xlsx", data_only=True)
ws = wb.active

rows = list(ws.iter_rows(values_only=True))
headers = list(rows[0])
data_rows = rows[1:]

# Build column index map (ignore None headers)
col = {name: idx for idx, name in enumerate(headers) if name is not None}

cleaned = []
ignored = []

name_missing_counter = 0
phone_missing_counter = 0

for row in data_rows:
    row = list(row)
    reasons = []

    name_val = row[col.get("Name")] if col.get("Name") is not None else None
    phone_val = row[col.get("Mobile No.")] if col.get("Mobile No.") is not None else None
    amount_val = row[col.get("Monthly Fees")] if col.get("Monthly Fees") is not None else None
    mode_val = row[col.get("Mode")] if col.get("Mode") is not None else None
    paydate_val = row[col.get("Payment date")] if col.get("Payment date") is not None else None

    name_missing = not name_val or str(name_val).strip() == ""
    phone_missing = not phone_val or str(phone_val).strip() == ""

    # Both name and phone missing → ignore
    if name_missing and phone_missing:
        reasons.append("missing both name and phone number")

    # Validate amount
    if not reasons and not is_valid_amount(amount_val):
        reasons.append("missing or invalid amount")

    # Validate payment date
    if not reasons:
        converted_date = format_date(paydate_val)
        if converted_date is None:
            reasons.append(f"missing or invalid payment date: {paydate_val!r}")

    # Validate mode
    if not reasons:
        normalized_mode = normalize_mode(mode_val)
        if normalized_mode is None:
            reasons.append(f"invalid payment frequency: {mode_val!r}")

    if reasons:
        ignored.append(row + ["; ".join(reasons)])
    else:
        # Fill placeholders
        if name_missing:
            name_missing_counter += 1
            row[col["Name"]] = f"NAME_MISSING_{name_missing_counter}"
        if phone_missing:
            phone_missing_counter += 1
            row[col["Mobile No."]] = f"PHONE_MISSING_{phone_missing_counter}"

        # Normalize mode
        row[col["Mode"]] = normalized_mode

        # Convert payment date to readable string
        row[col["Payment date"]] = converted_date

        # Convert Next Payment date too if it's a real date (skip formulas)
        np_idx = col.get("Next Payment")
        if np_idx is not None:
            np_val = row[np_idx]
            formatted_np = format_date(np_val)
            row[np_idx] = formatted_np if formatted_np else ""

        cleaned.append(row)

# Output only the meaningful headers (drop trailing None columns)
out_headers = [h for h in headers if h is not None]
num_cols = len(out_headers)

def format_row(row, num_cols):
    """Convert any datetime values to YYYY-MM-DD strings for CSV output."""
    result = []
    for val in row[:num_cols]:
        if isinstance(val, datetime):
            result.append(val.strftime("%Y-%m-%d"))
        elif isinstance(val, date):
            result.append(val.strftime("%Y-%m-%d"))
        else:
            result.append(val)
    return result


# Write cleaned.csv
with open("cleaned.csv", "w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)
    writer.writerow(out_headers)
    writer.writerows([format_row(r, num_cols) for r in cleaned])

# Write ignored.csv
with open("ignored.csv", "w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)
    writer.writerow(out_headers + ["ignore_reason"])
    writer.writerows([format_row(r, num_cols) + [r[-1]] for r in ignored])

print(f"Total data rows:         {len(data_rows)}")
print(f"Cleaned rows:            {len(cleaned)}")
print(f"Ignored rows:            {len(ignored)}")
print(f"Name placeholders used:  {name_missing_counter}")
print(f"Phone placeholders used: {phone_missing_counter}")
print(f"Total check:             {len(cleaned) + len(ignored)} (should be {len(data_rows)})")
print("Done. Files written: cleaned.csv, ignored.csv")
